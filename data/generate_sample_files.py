"""Generate PDF and DOCX sample resumes from text files for parser testing."""

from pathlib import Path

try:
    from fpdf import FPDF
    from docx import Document
except ImportError as exc:
    raise ImportError(
        "Please install dependencies: pip install fpdf2 python-docx"
    ) from exc


SAMPLE_RESUMES_DIR = Path(__file__).parent / "sample_resumes"


def txt_to_pdf(txt_path: Path, pdf_path: Path) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=11)
    text = txt_path.read_text(encoding="utf-8")
    # FPDF does not support unicode well; replace problematic chars
    text = text.encode("latin-1", errors="replace").decode("latin-1")
    for line in text.splitlines():
        pdf.cell(200, 6, txt=line, ln=True)
    pdf.output(str(pdf_path))


def txt_to_docx(txt_path: Path, docx_path: Path) -> None:
    doc = Document()
    text = txt_path.read_text(encoding="utf-8")
    for line in text.splitlines():
        doc.add_paragraph(line)
    doc.save(str(docx_path))


def main():
    txt_files = sorted(SAMPLE_RESUMES_DIR.glob("*.txt"))
    for txt_file in txt_files:
        base_name = txt_file.stem
        txt_to_pdf(txt_file, SAMPLE_RESUMES_DIR / f"{base_name}.pdf")
        txt_to_docx(txt_file, SAMPLE_RESUMES_DIR / f"{base_name}.docx")
        print(f"Generated PDF and DOCX for {txt_file.name}")


if __name__ == "__main__":
    main()
